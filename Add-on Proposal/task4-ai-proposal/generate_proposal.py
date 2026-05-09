from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import os

# ================= THEME =================
BG = RGBColor(15, 23, 42)        # dark navy
WHITE = RGBColor(255, 255, 255)
ACCENT = RGBColor(0, 255, 180)
GRAY = RGBColor(148, 163, 184)
CARD_BG = RGBColor(30, 41, 59)

# ================= HELPERS =================

def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

def text(shape, size=24, bold=False, color=WHITE):
    tf = shape.text_frame
    p = tf.paragraphs[0]
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color

def title(slide, txt):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(9), Inches(1))
    box.text = txt
    text(box, size=34, bold=True)
    return box

def card(slide, x, y, w, h, t, d):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = ACCENT

    tf = shape.text_frame
    tf.clear()

    p1 = tf.paragraphs[0]
    p1.text = t
    p1.font.size = Pt(16)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf.add_paragraph()
    p2.text = d
    p2.font.size = Pt(11)
    p2.font.color.rgb = GRAY

    return shape

# ================= CANVA-STYLE ANIMATION =================

def animated_steps(prs, builder, items):
    """
    Simulates Canva-like animation by progressive slides
    """
    for i in range(1, len(items) + 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        builder(slide, items[:i])

# ================= SECTION BREAK =================

def section(prs, txt):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(2))
    box.text = txt
    text(box, size=40, bold=True, color=ACCENT)

# ================= SLIDE BUILDERS =================

def feature_builder(slide, items):
    set_bg(slide)
    title(slide, "AI CHATBOT FEATURES")

    x, y = 0.6, 1.5
    for i, f in enumerate(items):
        card(slide, Inches(x), Inches(y), Inches(3), Inches(1.4),
             f, "Automated AI-powered system")
        x += 3.2
        if (i + 1) % 3 == 0:
            x = 0.6
            y += 1.8

def roi_builder(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "BUSINESS IMPACT")

    card(slide, Inches(1), Inches(2), Inches(8), Inches(3),
         "Expected ROI",
         "• 30% more leads\n• 90% faster replies\n• 50% less manual workload\n• Higher conversions")

# ================= MAIN =================

def build_pro_deck():
    prs = Presentation()

    # ========== COVER ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    t = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    t.text = "AI CHATBOT FOR LEAD GENERATION"
    text(t, size=42, bold=True)

    s = slide.shapes.add_textbox(Inches(1), Inches(3.3), Inches(8), Inches(1))
    s.text = "LA Digital Agency • Premium Automation Solutions"
    text(s, size=16, color=ACCENT)

    # ========== PROBLEM ==========
    section(prs, "PROBLEM STATEMENT")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "Key Issues")

    card(slide, Inches(0.7), Inches(1.6), Inches(4), Inches(2),
         "Slow Response", "Customers wait too long")

    card(slide, Inches(5), Inches(1.6), Inches(4), Inches(2),
         "Lost Leads", "No instant engagement")

    card(slide, Inches(0.7), Inches(4), Inches(4), Inches(2),
         "Manual Work", "High workload on staff")

    card(slide, Inches(5), Inches(4), Inches(4), Inches(2),
         "Low Conversion", "Delayed replies lose sales")

    # ========== SOLUTION ==========
    section(prs, "AI SOLUTION")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "24/7 AI Sales Assistant")

    card(slide, Inches(1), Inches(2), Inches(8), Inches(3),
         "Smart Automation Engine",
         "Instant replies • Lead capture • Qualification • CRM sync • WhatsApp integration")

    # ========== ANIMATED FEATURES ==========
    features = [
        "Instant AI Replies",
        "Smart Lead Capture",
        "WhatsApp Integration",
        "CRM Sync",
        "Auto Booking",
        "Lead Qualification"
    ]

    animated_steps(prs, feature_builder, features)

    # ========== ROI ==========
    roi_builder(prs)

    # ========== PRICING ==========
    section(prs, "PRICING")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    title(slide, "Investment Plan")

    card(slide, Inches(1), Inches(2), Inches(4), Inches(3),
         "Setup Fee", "PKR 80,000\nOne-time")

    card(slide, Inches(5), Inches(2), Inches(4), Inches(3),
         "Monthly", "PKR 25,000\nAI + Maintenance")

    # ========== FINAL ==========
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)

    t = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    t.text = "THANK YOU"
    text(t, size=44, bold=True, color=ACCENT)

    c = slide.shapes.add_textbox(Inches(1), Inches(3.5), Inches(8), Inches(1))
    c.text = "LA Digital Agency | info@ladigital.agency"
    text(c, size=16)

    prs.save("PRO_AI_PITCH_DECK.pptx")
    print("PRO deck generated successfully!")

if __name__ == "__main__":
    build_pro_deck()